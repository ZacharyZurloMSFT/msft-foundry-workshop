"""Document ingestion pipeline: extract, chunk, embed, index."""

import hashlib
import io
import logging
from typing import BinaryIO

from azure.search.documents import SearchClient
from azure.search.documents.models import IndexingResult

from app.config import settings
from app.models import SearchDocument

logger = logging.getLogger(__name__)

# Chunking defaults
DEFAULT_CHUNK_SIZE = 1000
DEFAULT_CHUNK_OVERLAP = 200

# Supported MIME/extension mapping
ALLOWED_EXTENSIONS: dict[str, str] = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
}

MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB


def extract_text(file: BinaryIO, filename: str) -> str:
    """Extract plain text from a PDF, DOCX, or TXT file."""
    lower = filename.lower()

    if lower.endswith(".pdf"):
        return _extract_pdf(file)
    elif lower.endswith(".docx"):
        return _extract_docx(file)
    elif lower.endswith(".txt"):
        return file.read().decode("utf-8", errors="replace")
    else:
        raise ValueError(f"Unsupported file type: {filename}")


def _extract_pdf(file: BinaryIO) -> str:
    from pypdf import PdfReader

    reader = PdfReader(file)
    pages: list[str] = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text)
    return "\n".join(pages)


def _extract_docx(file: BinaryIO) -> str:
    from docx import Document

    doc = Document(file)
    return "\n".join(p.text for p in doc.paragraphs if p.text)


def chunk_text(
    text: str,
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    chunk_overlap: int = DEFAULT_CHUNK_OVERLAP,
) -> list[str]:
    """Split text into fixed-size chunks with overlap."""
    if not text:
        return []

    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunks.append(text[start:end])
        start += chunk_size - chunk_overlap
    return chunks


def build_search_documents(
    filename: str,
    chunks: list[str],
    embeddings: list[list[float]],
) -> list[SearchDocument]:
    """Build SearchDocument objects for each chunk."""
    docs: list[SearchDocument] = []
    for idx, (chunk, embedding) in enumerate(zip(chunks, embeddings)):
        doc_id = hashlib.sha256(f"{filename}:{idx}".encode()).hexdigest()[:32]
        docs.append(
            SearchDocument(
                id=doc_id,
                content=chunk,
                content_vector=embedding,
                title=filename,
                source=filename,
                chunk_index=idx,
            )
        )
    return docs


def generate_embeddings(
    chunks: list[str],
    project_client: object | None = None,
) -> list[list[float]]:
    """Generate embeddings using Azure OpenAI via the project client or openai SDK."""
    if project_client is not None:
        return _embed_via_project_client(chunks, project_client)
    return _embed_via_openai_sdk(chunks)


def _embed_via_project_client(
    chunks: list[str],
    project_client: object,
) -> list[list[float]]:
    """Use AIProjectClient inference to get embeddings."""
    from azure.ai.projects import AIProjectClient

    client: AIProjectClient = project_client  # type: ignore[assignment]
    inference = client.inference
    response = inference.get_embeddings(
        model=settings.azure_openai_embedding_deployment,
        input=chunks,
    )
    return [item.embedding for item in response.data]


def _embed_via_openai_sdk(chunks: list[str]) -> list[list[float]]:
    """Fallback: use the openai SDK directly."""
    import openai

    response = openai.embeddings.create(
        model=settings.azure_openai_embedding_deployment,
        input=chunks,
    )
    return [item.embedding for item in response.data]


def index_documents(
    search_client: SearchClient,
    documents: list[SearchDocument],
) -> list[IndexingResult]:
    """Upload documents to Azure AI Search."""
    batch = [doc.model_dump() for doc in documents]
    result = search_client.upload_documents(documents=batch)
    logger.info("Indexed %d documents", len(batch))
    return result


def delete_documents_by_source(
    search_client: SearchClient,
    filename: str,
) -> int:
    """Delete all chunks for a given source filename from the index."""
    results = search_client.search(
        search_text="*",
        filter=f"source eq '{filename}'",
        select=["id"],
    )
    doc_ids = [{"id": r["id"]} for r in results]
    if doc_ids:
        search_client.delete_documents(documents=doc_ids)
        logger.info("Deleted %d chunks for '%s'", len(doc_ids), filename)
    return len(doc_ids)
